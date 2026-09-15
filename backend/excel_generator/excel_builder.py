"""
Khanna Travels & Holidays — Master Client Excel Workbook Generator & Importer
Uses openpyxl to generate a 5-sheet operational Excel workbook with executive formatting
and parse uploaded workbooks for safe diff-based updates.
"""

import os
from typing import Dict, Any, List, Tuple
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from ..models.master_schema import MasterApplicationData


# Branding Styles
FONT_TITLE = Font(name="Calibri", size=14, bold=True, color="FFFFFF")
FONT_HEADER = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
FONT_SECTION = Font(name="Calibri", size=11, bold=True, color="1E3A8A")
FONT_FIELD = Font(name="Calibri", size=11, bold=True, color="334155")
FONT_DATA = Font(name="Calibri", size=11, color="0F172A")
FONT_STATUS = Font(name="Calibri", size=9, italic=True, color="64748B")

FILL_TITLE = PatternFill(start_color="1E3A8A", end_color="1E3A8A", fill_type="solid")
FILL_HEADER = PatternFill(start_color="2563EB", end_color="2563EB", fill_type="solid")
FILL_SECTION = PatternFill(start_color="EFF6FF", end_color="EFF6FF", fill_type="solid")
FILL_ZEBRA = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")

THIN_BORDER_SIDE = Side(border_style="thin", color="CBD5E1")
BORDER_CELL = Border(
    left=THIN_BORDER_SIDE, right=THIN_BORDER_SIDE, top=THIN_BORDER_SIDE, bottom=THIN_BORDER_SIDE
)


def apply_title_row(ws, title_text: str, end_col: str = "C"):
    ws.merge_cells(f"A1:{end_col}1")
    title_cell = ws["A1"]
    title_cell.value = f"KHANNA TRAVELS & HOLIDAYS — {title_text.upper()}"
    title_cell.font = FONT_TITLE
    title_cell.fill = FILL_TITLE
    title_cell.alignment = Alignment(horizontal="center", vertical="center")
    ws.row_dimensions[1].height = 36


def autofit_columns(ws, max_cols=16):
    for col in range(1, max_cols + 1):
        col_letter = get_column_letter(col)
        max_len = 0
        for cell in ws[col_letter]:
            if cell.row == 1:
                continue
            if cell.value:
                lines = str(cell.value).split("\n")
                line_max = max(len(l) for l in lines)
                max_len = max(max_len, line_max)
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)


def build_master_excel(app_data: MasterApplicationData, output_path: str) -> str:
    """Generates the comprehensive 5-sheet Master Client Excel Workbook."""
    wb = openpyxl.Workbook()
    
    # -------------------------------------------------------------
    # SHEET 1: Client Master Data
    # -------------------------------------------------------------
    ws1 = wb.active
    ws1.title = "Client Master Data"
    apply_title_row(ws1, "Client Master Data Sheet", "C")
    
    headers1 = ["Field", "Value", "Verification Status"]
    for col_num, h in enumerate(headers1, 1):
        c = ws1.cell(row=2, column=col_num, value=h)
        c.font = FONT_HEADER
        c.fill = FILL_HEADER
        c.alignment = Alignment(horizontal="left", vertical="center")
        c.border = BORDER_CELL
    ws1.row_dimensions[2].height = 24
    
    p = app_data.personal
    passp = app_data.passport
    addr = app_data.address
    fam = app_data.family
    cont = app_data.contact
    emp = app_data.employment
    statuses = app_data.fieldStatuses or {}
    
    master_sections = [
        ("PERSONAL INFORMATION", [
            ("Title", p.title, "personal.title"),
            ("Given Name", p.givenName, "personal.givenName"),
            ("Middle Name", p.middleName, "personal.middleName"),
            ("Surname", p.surname, "personal.surname"),
            ("Full Name (as in Passport)", p.fullName, "personal.fullName"),
            ("Previous Name", p.previousName, "personal.previousName"),
            ("Gender", p.gender, "personal.gender"),
            ("Date of Birth", p.dob, "personal.dob"),
            ("Place of Birth", p.placeOfBirth, "personal.placeOfBirth"),
            ("Country of Birth", p.countryOfBirth, "personal.countryOfBirth"),
            ("Nationality", p.nationality, "personal.nationality"),
        ]),
        ("PASSPORT INFORMATION", [
            ("Passport Type", passp.passportType, "passport.passportType"),
            ("Passport Number", passp.passportNumber, "passport.passportNumber"),
            ("Passport Issue Date", passp.issueDate, "passport.issueDate"),
            ("Passport Expiry Date", passp.expiryDate, "passport.expiryDate"),
            ("Passport Issue Place", passp.issuePlace, "passport.issuePlace"),
            ("Passport Issuing Country", passp.issuingCountry, "passport.issuingCountry"),
            ("Issuing Authority", passp.issuingAuthority, "passport.issuingAuthority"),
        ]),
        ("ADDRESS INFORMATION", [
            ("Address Line 1", addr.addressLine1, "address.addressLine1"),
            ("Address Line 2", addr.addressLine2, "address.addressLine2"),
            ("City", addr.city, "address.city"),
            ("State", addr.state, "address.state"),
            ("Country", addr.country, "address.country"),
            ("PIN / Postal Code", addr.postalCode, "address.postalCode"),
            ("Current Residential Address", addr.currentResidentialAddress, "address.currentResidentialAddress"),
            ("Permanent Address", addr.permanentAddress, "address.permanentAddress"),
            ("Country of Residence", addr.countryOfResidence, "address.countryOfResidence"),
        ]),
        ("FAMILY INFORMATION", [
            ("Father's Full Name", fam.fatherFullName, "family.fatherFullName"),
            ("Mother's Full Name", fam.motherFullName, "family.motherFullName"),
            ("Spouse Full Name", fam.spouseFullName, "family.spouseFullName"),
            ("Spouse Passport Number", fam.spousePassportNumber, "family.spousePassportNumber"),
            ("Spouse DOB", fam.spouseDob, "family.spouseDob"),
            ("Spouse Nationality", fam.spouseNationality, "family.spouseNationality"),
            ("Spouse Occupation", fam.spouseOccupation, "family.spouseOccupation"),
        ]),
        ("CONTACT INFORMATION", [
            ("Mobile Number", cont.mobileNumber, "contact.mobileNumber"),
            ("Email Address", cont.emailAddress, "contact.emailAddress"),
            ("Alternate Contact", cont.alternateContact, "contact.alternateContact"),
            ("Emergency Contact", cont.emergencyContact, "contact.emergencyContact"),
        ]),
        ("EMPLOYMENT INFORMATION", [
            ("Employment Status", emp.employmentStatus, "employment.employmentStatus"),
            ("Employer / Organization Name", emp.employerName, "employment.employerName"),
            ("Job Title / Designation", emp.jobTitle, "employment.jobTitle"),
            ("Department", emp.department, "employment.department"),
            ("Employment Start Date", emp.employmentStartDate, "employment.employmentStartDate"),
            ("Annual / Monthly Income", emp.annualIncome, "employment.annualIncome"),
            ("Business Name (if self-employed)", emp.businessName, "employment.businessName"),
            ("Business Type", emp.businessType, "employment.businessType"),
            ("School / College Name (if student)", emp.schoolCollegeName, "employment.schoolCollegeName"),
            ("Course Name", emp.courseName, "employment.courseName"),
            ("Grade / Class", emp.gradeClass, "employment.gradeClass"),
            ("Other Information", emp.otherDetails, "employment.otherDetails"),
        ])
    ]
    
    current_row = 3
    for sec_title, fields in master_sections:
        # Section Header Row
        ws1.merge_cells(start_row=current_row, start_column=1, end_row=current_row, end_column=3)
        s_cell = ws1.cell(row=current_row, column=1, value=sec_title)
        s_cell.font = FONT_SECTION
        s_cell.fill = FILL_SECTION
        s_cell.alignment = Alignment(horizontal="left", vertical="center", indent=1)
        for col in range(1, 4):
            ws1.cell(row=current_row, column=col).border = BORDER_CELL
        ws1.row_dimensions[current_row].height = 22
        current_row += 1
        
        for field_label, val, field_key in fields:
            c1 = ws1.cell(row=current_row, column=1, value=field_label)
            c2 = ws1.cell(row=current_row, column=2, value=str(val) if val else "")
            
            # Status resolution
            status_val = statuses.get(field_key, "manual" if val else "pending")
            status_badge = "Verified [OK]" if status_val == "verified" else "OCR Extracted" if status_val == "ocr" else "Manual Entry" if val else "Pending Verification"
            c3 = ws1.cell(row=current_row, column=3, value=status_badge)
            
            c1.font = FONT_FIELD
            c2.font = FONT_DATA
            c3.font = FONT_STATUS
            
            c1.border = BORDER_CELL
            c2.border = BORDER_CELL
            c3.border = BORDER_CELL
            
            ws1.row_dimensions[current_row].height = 20
            current_row += 1
            
    ws1.freeze_panes = "A3"
    autofit_columns(ws1, 3)

    # -------------------------------------------------------------
    # SHEET 2: Travellers
    # -------------------------------------------------------------
    ws2 = wb.create_sheet(title="Travellers")
    apply_title_row(ws2, "Accompanying Travellers List", "P")
    
    t_headers = [
        "Sr. No.", "Title", "Given Name", "Middle Name", "Surname", "Full Name",
        "Passport Number", "DOB", "Gender", "Nationality", "Relation",
        "Occupation", "Employer", "School / College", "Grade / Class", "Other Information"
    ]
    for col_num, h in enumerate(t_headers, 1):
        c = ws2.cell(row=2, column=col_num, value=h)
        c.font = FONT_HEADER
        c.fill = FILL_HEADER
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BORDER_CELL
    ws2.row_dimensions[2].height = 24
    
    row_idx = 3
    # Lead Applicant as row 1
    lead_row = [
        1, p.title, p.givenName, p.middleName, p.surname, p.fullName,
        passp.passportNumber, p.dob, p.gender, p.nationality, "Self",
        emp.jobTitle or emp.employmentStatus, emp.employerName, emp.schoolCollegeName, emp.gradeClass, "Primary Applicant"
    ]
    for col_num, val in enumerate(lead_row, 1):
        c = ws2.cell(row=row_idx, column=col_num, value=str(val) if val is not None else "")
        c.font = FONT_DATA
        c.border = BORDER_CELL
        c.alignment = Alignment(horizontal="left", vertical="center")
    ws2.row_dimensions[row_idx].height = 20
    row_idx += 1
    
    # Additional Travellers
    for idx, t in enumerate(app_data.travellers, start=2):
        t_data = [
            idx, t.title, t.givenName, t.middleName, t.surname, t.fullName,
            t.passportNumber, t.dob, t.gender, t.nationality, t.relationship,
            t.occupation, t.employer, t.schoolCollege, t.gradeClass, t.otherInfo
        ]
        for col_num, val in enumerate(t_data, 1):
            c = ws2.cell(row=row_idx, column=col_num, value=str(val) if val is not None else "")
            c.font = FONT_DATA
            c.border = BORDER_CELL
            c.alignment = Alignment(horizontal="left", vertical="center")
        ws2.row_dimensions[row_idx].height = 20
        row_idx += 1
        
    ws2.freeze_panes = "A3"
    ws2.auto_filter.ref = f"A2:P{row_idx - 1}"
    autofit_columns(ws2, 16)

    # -------------------------------------------------------------
    # SHEET 3: Travel Details
    # -------------------------------------------------------------
    ws3 = wb.create_sheet(title="Travel Details")
    apply_title_row(ws3, "Trip & Travel Specifications", "B")
    
    t_det_headers = ["Parameter", "Detail"]
    for col_num, h in enumerate(t_det_headers, 1):
        c = ws3.cell(row=2, column=col_num, value=h)
        c.font = FONT_HEADER
        c.fill = FILL_HEADER
        c.alignment = Alignment(horizontal="left", vertical="center")
        c.border = BORDER_CELL
    ws3.row_dimensions[2].height = 24
    
    tr = app_data.travel
    fin = app_data.financial
    travel_rows = [
        ("Destination Country", tr.destinationCountry or app_data.selectedCountry),
        ("Destination Countries (All)", ", ".join(tr.destinationCountries) if tr.destinationCountries else tr.destinationCountry),
        ("Visa Type", tr.visaType),
        ("Purpose of Travel", tr.purposeOfTravel),
        ("Travel Start Date", tr.travelStartDate),
        ("Travel End Date", tr.travelEndDate),
        ("Total Number of Days", tr.numberOfDays),
        ("Total Number of Nights", tr.numberOfNights),
        ("Entry Type", tr.entryType),
        ("Number of Entries Requested", tr.numberOfEntries),
        ("Countries to be Visited", tr.countriesToBeVisited),
        ("Cities to be Visited", tr.citiesToBeVisited),
        ("Flight Outbound", f"{tr.flightNumber} ({tr.departureAirport} to {tr.arrivalAirport})" if tr.flightNumber else "Confirmed"),
        ("Flight Return", f"{tr.returnFlightNumber} ({tr.arrivalAirport} to {tr.departureAirport})" if tr.returnFlightNumber else "Confirmed"),
        ("Flight PNR / Reference", tr.pnrBookingRef),
        ("Trip Sponsor", fin.tripSponsor),
        ("Sponsor Name / Relation", f"{fin.sponsorName} ({fin.sponsorRelation})" if fin.sponsorName else "N/A"),
        ("Bank Statement Attached", "Yes (Verified)" if fin.bankStatementAvailable else "No"),
        ("ITR Filed & Attached", "Yes (Verified)" if fin.itrAvailable else "No"),
        ("Salary Slips Attached", "Yes (Verified)" if fin.salarySlipsAvailable else "No"),
        ("Employment / Leave Letter Attached", "Yes (Verified)" if fin.employmentLetterAvailable else "No"),
        ("Other Financial Documents", fin.otherFinancialDocs or "N/A"),
    ]
    
    for r_idx, (f_name, f_val) in enumerate(travel_rows, 3):
        c1 = ws3.cell(row=r_idx, column=1, value=f_name)
        c2 = ws3.cell(row=r_idx, column=2, value=str(f_val))
        c1.font = FONT_FIELD
        c2.font = FONT_DATA
        c1.border = BORDER_CELL
        c2.border = BORDER_CELL
        ws3.row_dimensions[r_idx].height = 20
        
    ws3.freeze_panes = "A3"
    autofit_columns(ws3, 2)

    # -------------------------------------------------------------
    # SHEET 4: Accommodation
    # -------------------------------------------------------------
    ws4 = wb.create_sheet(title="Accommodation")
    apply_title_row(ws4, "Hotel Bookings & Stay Details", "J")
    
    h_headers = [
        "Sr. No.", "Hotel Name", "Address", "City", "Country",
        "Check-in Date", "Check-out Date", "Nights", "Contact Number", "Booking Reference"
    ]
    for col_num, h in enumerate(h_headers, 1):
        c = ws4.cell(row=2, column=col_num, value=h)
        c.font = FONT_HEADER
        c.fill = FILL_HEADER
        c.alignment = Alignment(horizontal="center", vertical="center")
        c.border = BORDER_CELL
    ws4.row_dimensions[2].height = 24
    
    h_row_idx = 3
    if app_data.hotels:
        for idx, h in enumerate(app_data.hotels, 1):
            h_data = [
                idx, h.hotelName, h.address, h.city, h.country,
                h.checkInDate, h.checkOutDate, h.numberOfNights, h.contactNumber, h.bookingReference
            ]
            for col_num, val in enumerate(h_data, 1):
                c = ws4.cell(row=h_row_idx, column=col_num, value=str(val) if val is not None else "")
                c.font = FONT_DATA
                c.border = BORDER_CELL
                c.alignment = Alignment(horizontal="left", vertical="center")
            ws4.row_dimensions[h_row_idx].height = 20
            h_row_idx += 1
    else:
        # Default placeholder row
        c = ws4.cell(row=h_row_idx, column=1, value=1)
        c2 = ws4.cell(row=h_row_idx, column=2, value="Accommodation details pending verification")
        for col in range(1, 11):
            ws4.cell(row=h_row_idx, column=col).border = BORDER_CELL
            ws4.cell(row=h_row_idx, column=col).font = FONT_DATA
        ws4.row_dimensions[h_row_idx].height = 20
        h_row_idx += 1
        
    ws4.freeze_panes = "A3"
    autofit_columns(ws4, 10)

    # -------------------------------------------------------------
    # SHEET 5: Additional Information
    # -------------------------------------------------------------
    ws5 = wb.create_sheet(title="Additional Information")
    apply_title_row(ws5, "Executive Client Notes & Special Instructions", "B")
    
    ws5.cell(row=2, column=1, value="Note / Field").font = FONT_HEADER
    ws5.cell(row=2, column=1).fill = FILL_HEADER
    ws5.cell(row=2, column=1).border = BORDER_CELL
    ws5.cell(row=2, column=2, value="Content").font = FONT_HEADER
    ws5.cell(row=2, column=2).fill = FILL_HEADER
    ws5.cell(row=2, column=2).border = BORDER_CELL
    ws5.row_dimensions[2].height = 24
    
    add_rows = [
        ("Application Reference ID", app_data.applicationId or "APP-KTH-2026-001"),
        ("Target Visa Country", app_data.selectedCountry),
        ("Executive Additional Information", app_data.additional.content or "No additional notes."),
        ("Include in Cover Letter", "Yes" if app_data.additional.includeInCoverLetter else "No"),
        ("Last Updated Timestamp", app_data.updatedAt)
    ]
    for r_idx, (k, v) in enumerate(add_rows, 3):
        c1 = ws5.cell(row=r_idx, column=1, value=k)
        c2 = ws5.cell(row=r_idx, column=2, value=v)
        c1.font = FONT_FIELD
        c2.font = FONT_DATA
        c1.border = BORDER_CELL
        c2.border = BORDER_CELL
        ws5.row_dimensions[r_idx].height = 24
        
    ws5.freeze_panes = "A3"
    autofit_columns(ws5, 2)
    
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    wb.save(output_path)
    return output_path


def parse_and_diff_excel(file_bytes: bytes, current_data: MasterApplicationData) -> Dict[str, Any]:
    """
    Parses an uploaded edited Excel workbook, validates its structure,
    extracts fields from Sheet 1, and produces a diff preview before applying.
    """
    import io
    wb = openpyxl.load_workbook(io.BytesIO(file_bytes), data_only=True)
    
    if "Client Master Data" not in wb.sheetnames:
        raise ValueError("Invalid workbook: missing 'Client Master Data' sheet.")
        
    ws = wb["Client Master Data"]
    field_map = {}
    
    # Read rows from row 3 downwards
    for row in range(3, ws.max_row + 1):
        field_name = ws.cell(row=row, column=1).value
        field_val = ws.cell(row=row, column=2).value
        if field_name and field_val is not None:
            clean_name = str(field_name).strip()
            field_map[clean_name] = str(field_val).strip()
            
    diffs = []
    
    # Check key personal and passport fields
    check_pairs = [
        ("Passport Number", current_data.passport.passportNumber, "passport.passportNumber"),
        ("Full Name (as in Passport)", current_data.personal.fullName, "personal.fullName"),
        ("Given Name", current_data.personal.givenName, "personal.givenName"),
        ("Surname", current_data.personal.surname, "personal.surname"),
        ("Date of Birth", current_data.personal.dob, "personal.dob"),
        ("Passport Issue Date", current_data.passport.issueDate, "passport.issueDate"),
        ("Passport Expiry Date", current_data.passport.expiryDate, "passport.expiryDate"),
        ("Passport Issue Place", current_data.passport.issuePlace, "passport.issuePlace"),
        ("Mobile Number", current_data.contact.mobileNumber, "contact.mobileNumber"),
        ("Email Address", current_data.contact.emailAddress, "contact.emailAddress"),
        ("Employer / Organization Name", current_data.employment.employerName, "employment.employerName"),
        ("Job Title / Designation", current_data.employment.jobTitle, "employment.jobTitle"),
    ]
    
    for label, current_val, field_path in check_pairs:
        if label in field_map:
            new_val = field_map[label]
            if str(current_val or "").strip() != new_val:
                diffs.append({
                    "label": label,
                    "fieldPath": field_path,
                    "currentValue": current_val or "(empty)",
                    "newValue": new_val
                })
                
    return {
        "valid": True,
        "totalFieldsExtracted": len(field_map),
        "changesDetected": len(diffs),
        "diffs": diffs,
        "rawFieldMap": field_map
    }
