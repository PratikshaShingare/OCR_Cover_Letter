"""
Khanna Travels & Holidays — Multi-Sheet Professional Client Data Export Service
Generates an executive-styled multi-sheet workbook (Khanna_Travels_Client_Export.xlsx):
- Sheet 1: Passport Details
- Sheet 2: Travel Details
- Sheet 3: Contact Details
- Sheet 4: Employment & Financials

Applies corporate styling: Khanna Navy Blue headers (#1B365D), white bold text,
thin borders, zebra striping on alternating rows, auto-fit column widths,
and frozen header panes.
"""

import io
from typing import List, Optional
import openpyxl
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

from ..models.master_schema import MasterApplicationData


HEADER_FILL = PatternFill(start_color="1B365D", end_color="1B365D", fill_type="solid")
HEADER_FONT = Font(name="Calibri", size=11, bold=True, color="FFFFFF")
ZEBRA_FILL = PatternFill(start_color="F8FAFC", end_color="F8FAFC", fill_type="solid")
WHITE_FILL = PatternFill(start_color="FFFFFF", end_color="FFFFFF", fill_type="solid")

THIN_BORDER_SIDE = Side(border_style="thin", color="CBD5E1")
DATA_BORDER = Border(
    left=THIN_BORDER_SIDE,
    right=THIN_BORDER_SIDE,
    top=THIN_BORDER_SIDE,
    bottom=THIN_BORDER_SIDE
)

DATA_FONT = Font(name="Calibri", size=10, color="1E293B")
ALIGN_LEFT = Alignment(horizontal="left", vertical="center")
ALIGN_CENTER = Alignment(horizontal="center", vertical="center")


def apply_header_style(ws, headers: List[str]):
    ws.append(headers)
    for col_idx in range(1, len(headers) + 1):
        cell = ws.cell(row=1, column=col_idx)
        cell.fill = HEADER_FILL
        cell.font = HEADER_FONT
        cell.alignment = ALIGN_CENTER
        cell.border = DATA_BORDER
    ws.row_dimensions[1].height = 28
    ws.freeze_panes = "A2"


def style_data_row(ws, row_idx: int, values: List[str], is_zebra: bool = False):
    ws.append(values)
    fill = ZEBRA_FILL if is_zebra else WHITE_FILL
    for col_idx in range(1, len(values) + 1):
        cell = ws.cell(row=row_idx, column=col_idx)
        cell.font = DATA_FONT
        cell.fill = fill
        cell.border = DATA_BORDER
        val_str = str(values[col_idx - 1] or "")
        # Center dates, IDs, passport numbers, and phone numbers
        if any(char.isdigit() for char in val_str) and len(val_str) < 18:
            cell.alignment = ALIGN_CENTER
        else:
            cell.alignment = ALIGN_LEFT
    ws.row_dimensions[row_idx].height = 22


def auto_fit_columns(ws):
    for col in ws.columns:
        max_len = 0
        col_letter = get_column_letter(col[0].column)
        for cell in col:
            val = str(cell.value or "")
            max_len = max(max_len, len(val))
        ws.column_dimensions[col_letter].width = max(max_len + 4, 14)


def build_multi_sheet_workbook(applications: List[MasterApplicationData]) -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    # Remove default sheet
    wb.remove(wb.active)

    # -------------------------------------------------------------
    # Sheet 1: Passport Details
    # -------------------------------------------------------------
    ws_pass = wb.create_sheet(title="Passport Details")
    pass_headers = [
        "Application ID", "Traveller Type", "Relationship", "Title", "Full Name",
        "Given Name", "Surname", "Passport Number", "Nationality", "Date of Birth",
        "Gender", "Place of Birth", "Place of Issue", "Issue Date", "Expiry Date",
        "Father Full Name", "Mother Full Name", "Spouse Full Name"
    ]
    apply_header_style(ws_pass, pass_headers)

    pass_row = 2
    for app in applications:
        p = app.personal
        pass_info = app.passport
        fam = app.family
        # Primary applicant row
        primary_vals = [
            app.applicationId, "Primary Applicant", "Self", p.title or "Mr.",
            p.fullName or f"{p.givenName} {p.surname}".strip(),
            p.givenName, p.surname, pass_info.passportNumber,
            p.nationality or "Indian", p.dob, p.gender, p.placeOfBirth,
            pass_info.issuePlace, pass_info.issueDate, pass_info.expiryDate,
            fam.fatherFullName, fam.motherFullName, fam.spouseFullName
        ]
        style_data_row(ws_pass, pass_row, primary_vals, is_zebra=(pass_row % 2 == 0))
        pass_row += 1

        # Accompanying travellers
        for t in app.travellers:
            t_vals = [
                app.applicationId, "Accompanying", t.relationship or "Companion",
                t.title or "Mr.", t.fullName or f"{t.givenName} {t.surname}".strip(),
                t.givenName, t.surname, t.passportNumber,
                t.nationality or "Indian", t.dob, t.gender, "",
                t.issuePlace, t.issueDate, t.expiryDate,
                "", "", ""
            ]
            style_data_row(ws_pass, pass_row, t_vals, is_zebra=(pass_row % 2 == 0))
            pass_row += 1

    auto_fit_columns(ws_pass)

    # -------------------------------------------------------------
    # Sheet 2: Travel Details
    # -------------------------------------------------------------
    ws_trav = wb.create_sheet(title="Travel Details")
    trav_headers = [
        "Application ID", "Destination Country", "Visa Type", "Travel Start Date",
        "Travel End Date", "Duration (Days)", "Purpose of Visit", "Flight Number",
        "Hotel Name", "Hotel Address", "Hotel City"
    ]
    apply_header_style(ws_trav, trav_headers)

    trav_row = 2
    for app in applications:
        tr = app.travel
        hotel_names = ", ".join([h.hotelName for h in app.hotels if h.hotelName]) if app.hotels else ""
        hotel_addrs = ", ".join([h.address for h in app.hotels if h.address]) if app.hotels else ""
        hotel_cities = ", ".join([h.city for h in app.hotels if h.city]) if app.hotels else ""

        # Calculate duration if dates available
        dur_days = ""
        try:
            from datetime import datetime
            if tr.travelStartDate and tr.travelEndDate:
                d1 = datetime.strptime(tr.travelStartDate, "%Y-%m-%d")
                d2 = datetime.strptime(tr.travelEndDate, "%Y-%m-%d")
                dur_days = str(max(1, (d2 - d1).days))
            elif tr.numberOfDays:
                dur_days = str(tr.numberOfDays)
        except Exception:
            dur_days = ""

        trav_vals = [
            app.applicationId, app.selectedCountry or "Standard", tr.visaType or "Tourist",
            tr.travelStartDate, tr.travelEndDate, dur_days,
            tr.purposeOfTravel or "Tourism / Vacation",
            tr.flightNumber or "",
            hotel_names, hotel_addrs, hotel_cities
        ]
        style_data_row(ws_trav, trav_row, trav_vals, is_zebra=(trav_row % 2 == 0))
        trav_row += 1

    auto_fit_columns(ws_trav)

    # -------------------------------------------------------------
    # Sheet 3: Contact Details
    # -------------------------------------------------------------
    ws_cont = wb.create_sheet(title="Contact Details")
    cont_headers = [
        "Application ID", "Full Name", "Country Code", "Mobile Number", "Email Address",
        "Alternate Contact", "Emergency Contact", "Address Line 1", "Address Line 2",
        "City", "State", "PIN / Postal Code", "Country"
    ]
    apply_header_style(ws_cont, cont_headers)

    cont_row = 2
    for app in applications:
        p = app.personal
        c = app.contact
        a = app.address
        cont_vals = [
            app.applicationId, p.fullName or f"{p.givenName} {p.surname}".strip(),
            getattr(c, "countryCode", "+91") or "+91", c.mobileNumber, c.emailAddress,
            c.alternateContact, c.emergencyContact,
            a.addressLine1 or a.currentResidentialAddress, a.addressLine2,
            a.city, a.state, a.postalCode, a.country or "India"
        ]
        style_data_row(ws_cont, cont_row, cont_vals, is_zebra=(cont_row % 2 == 0))
        cont_row += 1

    auto_fit_columns(ws_cont)

    # -------------------------------------------------------------
    # Sheet 4: Employment & Financials
    # -------------------------------------------------------------
    ws_emp = wb.create_sheet(title="Employment & Financials")
    emp_headers = [
        "Application ID", "Full Name", "Employment Status", "Employer / Organization",
        "Job Title / Role", "Department", "Start Date", "Annual Income",
        "Business Name", "Business Type", "School / College", "Course / Degree", "Grade / Year"
    ]
    apply_header_style(ws_emp, emp_headers)

    emp_row = 2
    for app in applications:
        p = app.personal
        e = app.employment
        emp_vals = [
            app.applicationId, p.fullName or f"{p.givenName} {p.surname}".strip(),
            e.employmentStatus or "Employed", e.employerName, e.jobTitle,
            e.department, e.employmentStartDate, e.annualIncome,
            e.businessName, e.businessType, e.schoolCollegeName,
            e.courseName, e.gradeClass
        ]
        style_data_row(ws_emp, emp_row, emp_vals, is_zebra=(emp_row % 2 == 0))
        emp_row += 1

    auto_fit_columns(ws_emp)

    return wb


def export_applications_to_bytes(applications: List[MasterApplicationData]) -> io.BytesIO:
    wb = build_multi_sheet_workbook(applications)
    buffer = io.BytesIO()
    wb.save(buffer)
    buffer.seek(0)
    return buffer
