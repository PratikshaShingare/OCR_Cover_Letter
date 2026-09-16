"""
Unit tests for Khanna Travels & Holidays Visa Automation System
Tests for new Phase 1-8 enhancements:
- Passport Normalizer
- Structured Address & Title Extraction
- Multi-sheet Excel Generator (4-Sheet corporate workbook)
- Admin Bulk Data Ingestion & Audit Log
"""

import os
import io
import csv
import unittest
from fastapi.testclient import TestClient
from PIL import Image

from backend.main import app
from backend.models.master_schema import MasterApplicationData, Traveller
from backend.services.passport_normalizer import normalize_image_buffer, auto_trim_margins
from backend.ocr.passport_extractor import parse_structured_address, extract_document_data
from backend.services.multi_sheet_excel import export_applications_to_bytes, build_multi_sheet_workbook

client = TestClient(app)

class TestEnhancements(unittest.TestCase):

    def test_passport_normalizer_image(self):
        # Create small test image
        img = Image.new("RGB", (200, 300), color="white")
        buf = io.BytesIO()
        img.save(buf, format="JPEG")
        raw_bytes = buf.getvalue()

        normalized_bytes, mime = normalize_image_buffer(raw_bytes)
        self.assertEqual(mime, "image/jpeg")
        self.assertGreater(len(normalized_bytes), 0)

        # Open back with PIL to verify valid JPEG
        res_img = Image.open(io.BytesIO(normalized_bytes))
        self.assertGreater(res_img.width, 0)
        self.assertGreater(res_img.height, 0)

    def test_structured_address_parsing(self):
        sample_addr = "B-402, PALM SPRINGS, LINK ROAD, ANDHERI WEST, MUMBAI, MAHARASHTRA 400053 INDIA"
        parsed = parse_structured_address(sample_addr)

        self.assertTrue(bool(parsed.get("addressLine1")))
        self.assertEqual(parsed.get("country"), "India")
        self.assertEqual(parsed.get("postalCode"), "400053")
        self.assertTrue("MAHARASHTRA" in parsed.get("state", "").upper() or "MUMBAI" in parsed.get("city", "").upper())

    def test_multi_sheet_excel_single_client(self):
        app_data = MasterApplicationData(applicationId="APP-TEST-001")
        app_data.personal.fullName = "Rajesh Sharma"
        app_data.personal.gender = "Male"
        app_data.personal.nationality = "Indian"
        app_data.passport.passportNumber = "Z9876543"
        app_data.contact.mobileNumber = "+91 9876543210"
        app_data.contact.emailAddress = "rajesh@example.com"
        app_data.travel.destinationCountry = "Japan"
        app_data.travel.visaType = "Tourism"
        app_data.employment.employerName = "Tata Consultancy Services"
        app_data.employment.jobTitle = "Senior Lead Consultant"

        # Add a traveller
        t = Traveller(fullName="Pooja Sharma", relationship="Spouse", passportNumber="Z9876544", nationality="Indian")
        app_data.travellers.append(t)

        excel_buf = export_applications_to_bytes([app_data])
        excel_bytes = excel_buf.getvalue()
        self.assertGreater(len(excel_bytes), 1000)

        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        expected_sheets = ["Passport Details", "Travel Details", "Contact Details", "Employment & Financials"]
        for sheet_name in expected_sheets:
            self.assertIn(sheet_name, wb.sheetnames)

        # Verify header on Passport Details
        ws_passport = wb["Passport Details"]
        self.assertEqual(ws_passport.cell(row=1, column=1).value, "Application ID")

    def test_multi_sheet_excel_all_clients(self):
        app1 = MasterApplicationData(applicationId="APP-ALL-001")
        app1.personal.fullName = "Amit Patel"
        app1.passport.passportNumber = "A1234567"

        app2 = MasterApplicationData(applicationId="APP-ALL-002")
        app2.personal.fullName = "Sneha Patel"
        app2.passport.passportNumber = "A1234568"

        excel_buf = export_applications_to_bytes([app1, app2])
        excel_bytes = excel_buf.getvalue()
        self.assertGreater(len(excel_bytes), 1000)

        import openpyxl
        wb = openpyxl.load_workbook(io.BytesIO(excel_bytes))
        self.assertIn("Passport Details", wb.sheetnames)
        ws = wb["Passport Details"]
        # Header + 2 rows = row 3
        self.assertGreaterEqual(ws.max_row, 3)

    def test_admin_import_flow(self):
        # 1. Prepare sample CSV
        csv_buf = io.StringIO()
        writer = csv.writer(csv_buf)
        writer.writerow(["Passport No", "Full Name", "Gender", "Destination", "Mobile"])
        writer.writerow(["T1122334", "Vikram Malhotra", "Male", "France", "9820012345"])
        writer.writerow(["T1122335", "Sunita Malhotra", "Female", "France", "9820012346"])
        csv_bytes = csv_buf.getvalue().encode("utf-8")

        # 2. Upload file to /api/admin/import/upload
        res = client.post(
            "/api/admin/import/upload",
            files={"file": ("test_bulk.csv", csv_bytes, "text/csv")}
        )
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertTrue(data["success"])
        temp_id = data["tempId"]
        self.assertEqual(data["totalRows"], 2)
        self.assertIn("detectedMappings", data)

        # 3. Execute bulk import
        mappings = {
            "Passport No": "passport.passportNumber",
            "Full Name": "personal.fullName",
            "Gender": "personal.gender",
            "Destination": "selectedCountry",
            "Mobile": "contact.mobileNumber"
        }
        exec_res = client.post(
            "/api/admin/import/execute",
            json={
                "tempId": temp_id,
                "mappings": mappings,
                "duplicateStrategy": "update"
            }
        )
        self.assertEqual(exec_res.status_code, 200)
        exec_data = exec_res.json()
        self.assertTrue(exec_data["success"])
        self.assertEqual(exec_data["importedCount"] + exec_data["updatedCount"], 2)

        # 4. Check import history
        hist_res = client.get("/api/admin/import/history")
        self.assertEqual(hist_res.status_code, 200)
        hist_data = hist_res.json()
        self.assertIsInstance(hist_data, list)
        self.assertGreater(len(hist_data), 0)
        self.assertEqual(hist_data[0]["filename"], "test_bulk.csv")

if __name__ == "__main__":
    unittest.main()
