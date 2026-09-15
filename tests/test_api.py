"""
FastAPI End-to-End API Integration Tests
Khanna Travels & Holidays — Visa Document Automation System
"""

import sys
import os
import io
import unittest
from fastapi.testclient import TestClient

sys.path.insert(0, os.path.abspath(os.path.join(os.path.dirname(__file__), "..")))
from backend.main import app

client = TestClient(app)


class TestApiIntegration(unittest.TestCase):
    
    def test_01_health(self):
        res = client.get("/api/health")
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.json()["status"], "healthy")

    def test_02_templates_list(self):
        res = client.get("/api/templates")
        self.assertEqual(res.status_code, 200)
        tpls = res.json()
        self.assertEqual(len(tpls), 3)
        ids = [t["id"] for t in tpls]
        self.assertIn("standard", ids)
        self.assertIn("japan", ids)
        self.assertIn("singapore", ids)

    def test_03_application_save_and_internal_excel(self):
        """Verify saving an application creates record and auto-generates internal Excel."""
        payload = {
            "applicationId": "APP-TEST-API-001",
            "selectedCountry": "France",
            "personal": {
                "fullName": "Mr. Rahul Sharma",
                "givenName": "Rahul",
                "surname": "Sharma",
                "nationality": "Indian"
            },
            "passport": {
                "passportNumber": "T9988776",
                "issuePlace": "Mumbai",
                "issueDate": "2024-01-10",
                "expiryDate": "2034-01-09"
            },
            "travel": {
                "destinationCountry": "France",
                "travelStartDate": "2026-11-01",
                "travelEndDate": "2026-11-15",
                "purposeOfTravel": "Tourism"
            }
        }
        res = client.post("/api/applications", json=payload)
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertEqual(data["applicationId"], "APP-TEST-API-001")
        self.assertEqual(data["personal"]["fullName"], "Mr. Rahul Sharma")

        # Verify Master Excel workbook was automatically updated on disk
        import openpyxl
        master_excel_path = os.path.join("data", "Khanna_Travels_Client_Master.xlsx")
        self.assertTrue(os.path.exists(master_excel_path))
        wb = openpyxl.load_workbook(master_excel_path)
        self.assertEqual(len(wb.sheetnames), 1)
        self.assertEqual(wb.sheetnames[0], "Client Data")
        ws = wb["Client Data"]
        found = any(ws.cell(row=r, column=1).value == "APP-TEST-API-001" for r in range(2, ws.max_row + 1))
        self.assertTrue(found)
        wb.close()

    def test_04_application_list_and_get(self):
        """Verify listing and fetching single application."""
        list_res = client.get("/api/applications")
        self.assertEqual(list_res.status_code, 200)
        apps = list_res.json()
        self.assertGreaterEqual(len(apps), 1)
        found = any(a.get("id") == "APP-TEST-API-001" for a in apps)
        self.assertTrue(found)

        get_res = client.get("/api/applications/APP-TEST-API-001")
        self.assertEqual(get_res.status_code, 200)
        app_data = get_res.json()
        self.assertEqual(app_data["passport"]["passportNumber"], "T9988776")

    def test_05_generate_cover_letter_api(self):
        """Verify cover letter generation API endpoint with clean data."""
        payload = {
            "applicationId": "APP-TEST-GEN-01",
            "selectedCountry": "Japan",
            "personal": {
                "fullName": "Ms. Priya Mehta",
                "nationality": "Indian"
            },
            "passport": {
                "passportNumber": "J1234567"
            },
            "travel": {
                "destinationCountry": "Japan",
                "purposeOfTravel": "Holiday"
            },
            "hotels": [
                {
                    "hotelName": "Tokyo Hotel",
                    "city": "Tokyo",
                    "checkInDate": "2026-07-01",
                    "checkOutDate": "2026-07-07",
                    "contactNumber": "+81 3-1111-2222"
                }
            ]
        }
        gen_res = client.post("/api/documents/generate-cover-letter", json=payload)
        self.assertEqual(gen_res.status_code, 200)
        doc = gen_res.json()["document"]
        self.assertIn("html", doc)
        self.assertIn("Consulate General of Japan", doc["html"])
        self.assertIn("Priya Mehta", doc["html"])

    def test_06_download_docx_api(self):
        """Verify DOCX binary download."""
        payload = {
            "selectedCountry": "France",
            "personal": {
                "fullName": "Mr. Rahul Sharma"
            },
            "passport": {
                "passportNumber": "T9988776"
            },
            "travel": {
                "destinationCountry": "France"
            }
        }
        res = client.post("/api/documents/download-docx", json=payload)
        self.assertEqual(res.status_code, 200)
        self.assertIn("wordprocessingml", res.headers.get("content-type", ""))
        self.assertGreater(len(res.content), 2000)

    def test_07_download_pdf_api(self):
        """Verify PDF binary download."""
        payload = {
            "selectedCountry": "Singapore",
            "personal": {
                "fullName": "Mr. Vikram Patel"
            },
            "passport": {
                "passportNumber": "P4567890"
            },
            "travel": {
                "destinationCountry": "Singapore"
            }
        }
        res = client.post("/api/documents/download-pdf", json=payload)
        self.assertEqual(res.status_code, 200)
        self.assertEqual(res.headers.get("content-type"), "application/pdf")
        self.assertGreater(len(res.content), 2000)

    def test_08_upload_passport_ocr(self):
        """Verify passport file upload OCR endpoint structure."""
        dummy_content = b"%PDF-1.4 dummy pdf content for OCR testing"
        files = {"file": ("test_passport.pdf", io.BytesIO(dummy_content), "application/pdf")}
        res = client.post("/api/ocr/upload-passport", files=files)
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertIn("passportNumber", data)
        self.assertIn("confidence", data)
        self.assertIn("fieldStatuses", data)

    def test_09_create_new_blank_application(self):
        """Verify POST /api/applications/new creates completely blank record with sequential ID."""
        res = client.post("/api/applications/new")
        self.assertEqual(res.status_code, 200)
        data = res.json()
        self.assertTrue(data["applicationId"].startswith("APP-2026-"))
        self.assertEqual(data["personal"]["fullName"], "")
        self.assertEqual(data["passport"]["passportNumber"], "")

    def test_10_upload_and_get_passport_file(self):
        """Verify uploading passport document and streaming it back for preview."""
        # Create app
        app_res = client.post("/api/applications/new")
        app_id = app_res.json()["applicationId"]

        # Upload dummy PDF
        dummy_pdf = b"%PDF-1.4 %fake pdf binary for testing preview"
        files = {"file": ("passport_scan.pdf", io.BytesIO(dummy_pdf), "application/pdf")}
        upload_res = client.post(f"/api/applications/{app_id}/passport", files=files)
        self.assertEqual(upload_res.status_code, 200)
        upload_data = upload_res.json()
        self.assertTrue(upload_data["fileInfo"]["isPdf"])

        # Stream back preview
        preview_res = client.get(f"/api/applications/{app_id}/passport-file")
        self.assertEqual(preview_res.status_code, 200)
        self.assertEqual(preview_res.headers.get("content-type"), "application/pdf")
        self.assertEqual(preview_res.content, dummy_pdf)
    def test_11_delete_application(self):
        """Verify deleting an application completely purges record and files."""
        app_res = client.post("/api/applications/new")
        app_id = app_res.json()["applicationId"]

        del_res = client.delete(f"/api/applications/{app_id}")
        self.assertEqual(del_res.status_code, 200)
        self.assertTrue(del_res.json()["success"])

        get_res = client.get(f"/api/applications/{app_id}")
        self.assertEqual(get_res.status_code, 404)

    def test_12_passport_preview_and_info_endpoints(self):
        """Verify passport preview endpoint and passport info endpoint."""
        app_res = client.post("/api/applications/new")
        app_id = app_res.json()["applicationId"]

        # Check passport info before upload
        info_res = client.get(f"/api/applications/{app_id}/passport-info")
        self.assertEqual(info_res.status_code, 200)
        self.assertFalse(info_res.json()["hasPassport"])

        # Upload valid sample PDF if available
        user_uploaded_pdf = os.path.join(
            os.path.expanduser("~"),
            ".gemini", "antigravity", "brain",
            "f9e4d87f-c5c8-4340-9fb9-d57dc89259c3",
            ".user_uploaded", "media_1789196443697.pdf"
        )
        if os.path.exists(user_uploaded_pdf):
            with open(user_uploaded_pdf, "rb") as f:
                up_res = client.post(f"/api/applications/{app_id}/passport", files={"file": ("sample.pdf", f, "application/pdf")})
            self.assertEqual(up_res.status_code, 200)
            up_data = up_res.json()
            self.assertEqual(up_data["extracted"]["passportNumber"], "W4832498")

            # Check preview endpoint
            prev_res = client.get(f"/api/applications/{app_id}/passport-preview?page=1")
            self.assertEqual(prev_res.status_code, 200)
            self.assertEqual(prev_res.headers.get("content-type"), "image/png")
            self.assertTrue(prev_res.content.startswith(b"\x89PNG\r\n\x1a\n"))

            # Check passport info endpoint
            info_res = client.get(f"/api/applications/{app_id}/passport-info")
            self.assertEqual(info_res.status_code, 200)
            self.assertTrue(info_res.json()["hasPassport"])
            self.assertGreaterEqual(info_res.json()["totalPages"], 1)

        # Cleanup
        client.delete(f"/api/applications/{app_id}")


if __name__ == "__main__":
    unittest.main()

